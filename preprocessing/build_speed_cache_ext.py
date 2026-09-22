"""
Extend the traffic SPEED dataset to 2026-07, mirroring build_volume_cache_ext.py's fix for the
exact same staleness issue: smhan's pipeline_nowcast/speed_hourly_cache.npz stops at 2025-12-31
even though /path/to/raw_data/topis_speed/*.xlsx already has raw files through 2026_07
(same "already downloaded but never re-aggregated" gap volume had). Parses the 7 missing months
directly from the raw xlsx, restricted to the SAME 396 link_ids already in speed_hourly_cache.npz
(physical sensors -- set doesn't change), and concatenates onto the existing cache. Writes to
THIS project's own directory -- never touches smhan's original cache or raw files.
"""
import numpy as np
import openpyxl

GTS = "/home/ncrc/work/gts"
TOPIS_DIR = "/path/to/raw_data/topis_speed"
SRC_CACHE = "/path/to/raw_data/pipeline_nowcast/speed_hourly_cache.npz"
EXT_MONTHS = [(2026, m) for m in range(1, 8)]  # Jan..Jul 2026

orig = np.load(SRC_CACHE, allow_pickle=True)
known_ids = set(orig["link_ids"].tolist())
print(f"existing cache: {len(known_ids)} links, dates {orig['dates'].min()}..{orig['dates'].max()}, "
      f"{orig['speed'].shape}")

# accumulate (date, hour) -> per-link speed value, restricted to known link ids
new_by_dh = {}  # (date,hour) -> dict(link_id -> speed)
for (y, m) in EXT_MONTHS:
    path = f"{TOPIS_DIR}/{y}_{m:02d}.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    n_rows = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        date = row[0]
        link_id = str(row[3])
        if link_id not in known_ids:
            continue
        hourly = row[12:36]
        for h, v in enumerate(hourly):
            if v is None:
                continue
            new_by_dh.setdefault((date, h), {})[link_id] = float(v)
        n_rows += 1
    wb.close()
    print(f"  {y}-{m:02d}: {n_rows} matching rows")

# build new (dates, hours, speed) arrays in the SAME layout as the original cache: one column per
# (date,hour), speed[i, col] = value for link_ids[i] (NaN if missing that link/date/hour)
link_ids = list(orig["link_ids"])
link_pos = {lid: i for i, lid in enumerate(link_ids)}
dh_keys = sorted(new_by_dh.keys())
n_new_cols = len(dh_keys)
new_dates = np.array([k[0] for k in dh_keys], dtype=np.int64)
new_hours = np.array([k[1] for k in dh_keys], dtype=np.int64)
new_speed = np.full((len(link_ids), n_new_cols), np.nan, dtype=np.float32)
for c, (date, h) in enumerate(dh_keys):
    for lid, v in new_by_dh[(date, h)].items():
        new_speed[link_pos[lid], c] = v

dates = np.concatenate([orig["dates"], new_dates])
hours = np.concatenate([orig["hours"], new_hours])
speed = np.concatenate([orig["speed"], new_speed], axis=1)

np.savez(f"{GTS}/speed_hourly_cache_ext.npz",
         link_ids=orig["link_ids"], rows=orig["rows"], cols=orig["cols"],
         dates=dates, hours=hours, speed=speed)
print(f"saved speed_hourly_cache_ext.npz: dates {dates.min()}..{dates.max()}, speed {speed.shape}, "
      f"added {n_new_cols} new (date,hour) columns")
